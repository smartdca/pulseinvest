"""
市場溫度計：資料來源實測腳本
每次執行會逐一抓取 8 個來源，把成功/失敗、數值、失敗原因寫進 Google Sheets 的 ThermoTest 分頁。
只做測試，不影響網站任何功能。
"""

import io
import json
import os
import re
import sys
import time
import html as htmllib
from datetime import datetime
from zoneinfo import ZoneInfo

import requests

SHEET_ID = "1msqXlcSfVVNkfPFzMQelHQq3YjDrkwP5B95liXylNug"
TAB = "ThermoTest"
TIMEOUT = 25
HEADER = ["時間(台北)", "來源", "結果", "數值", "說明", "耗時(秒)", "失敗原因"]

BROWSER = {
    "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 "
                  "(KHTML, like Gecko) Chrome/126.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "en-US,en;q=0.9",
}


class FetchError(Exception):
    def __init__(self, reason):
        super().__init__(reason)
        self.reason = reason


# ---------- 共用工具 ----------

def get(url, extra_headers=None):
    headers = dict(BROWSER)
    if extra_headers:
        headers.update(extra_headers)
    try:
        r = requests.get(url, headers=headers, timeout=TIMEOUT)
    except requests.Timeout:
        raise FetchError("逾時")
    except requests.RequestException as e:
        raise FetchError(f"連線失敗 ({type(e).__name__})")
    if r.status_code in (401, 403, 429, 503):
        raise FetchError(f"被擋 (HTTP {r.status_code})")
    if r.status_code != 200:
        raise FetchError(f"HTTP {r.status_code}")
    head = r.text[:4000].lower()
    if "<title>just a moment" in head or "<title>attention required" in head:
        raise FetchError("被擋 (防爬蟲驗證頁)")
    return r


def page_text(raw):
    t = re.sub(r"(?is)<(script|style|noscript)[^>]*>.*?</\1>", " ", raw)
    t = re.sub(r"(?s)<[^>]+>", " ", t)
    t = htmllib.unescape(t)
    return re.sub(r"\s+", " ", t)


def search(patterns, raw):
    """先在去掉標籤的文字裡找，找不到再到原始 HTML 裡找（有些數值藏在內嵌資料中）"""
    for source in (page_text(raw), htmllib.unescape(raw)):
        for p in patterns:
            m = re.search(p, source, re.I | re.S)
            if m:
                s = max(m.start() - 30, 0)
                snippet = re.sub(r"\s+", " ", source[s:m.end() + 30]).strip()
                return m, snippet[:120]
    return None, ""


def check_range(v, lo, hi):
    if not (lo <= v <= hi):
        raise FetchError(f"數值超出合理範圍 ({v})")
    return v


# ---------- 各來源 ----------

def src_cnn():
    r = get("https://production.dataviz.cnn.io/index/fearandgreed/graphdata",
            {"Accept": "application/json",
             "Referer": "https://www.cnn.com/",
             "Origin": "https://www.cnn.com"})
    try:
        fg = r.json()["fear_and_greed"]
        v = round(float(fg["score"]), 1)
    except Exception:
        raise FetchError("格式不符 (可能改版)")
    return check_range(v, 0, 100), str(fg.get("rating", ""))


def src_thetrading():
    r = get("https://www.thetrading.tools/fear-greed-index")
    m, snip = search([r"reads\s*(\d{1,3}(?:\.\d+)?)\s*/\s*100",
                      r"Greed Index[^0-9]{0,60}?(\d{1,3})\s*/\s*100"], r.text)
    if not m:
        raise FetchError("找不到數值 (可能是動態載入或改版)")
    return check_range(float(m.group(1)), 0, 100), snip


def src_cfgi():
    r = get("https://cfgi.io/stock-market-fear-and-greed-index/")
    m, snip = search([r"Fear and Greed Index Today:\s*(\d{1,3})",
                      r"stock market Fear and Greed Index today is\s*(\d{1,3})"], r.text)
    if not m:
        raise FetchError("找不到數值 (可能是動態載入或改版)")
    return check_range(float(m.group(1)), 0, 100), snip


def src_altindex():
    r = get("https://altindex.com/fear-and-greed-index")
    m, snip = search([r"Today.{0,3}s reading:?\s*(\d{1,3})",
                      r"Fear\s*&\s*Greed Index[^0-9]{0,80}?(\d{1,3})\s*[—-]\s*(?:Extreme )?(?:Fear|Greed|Neutral)"],
                     r.text)
    if not m:
        raise FetchError("找不到數值 (可能是動態載入或改版)")
    return check_range(float(m.group(1)), 0, 100), snip


def src_sentisense():
    last_err = "找不到數值"
    for url in ("https://sentisense.ai/market-mood", "https://sentisense.ai/"):
        try:
            r = get(url)
        except FetchError as e:
            last_err = e.reason
            continue
        m, snip = search([r"Market Mood[^0-9]{0,40}?(\d{1,3})\s*[·\-—|]?\s*(?:Extreme Fear|Extreme Greed|Fear|Greed|Neutral)"],
                         r.text)
        if m:
            note = f"{url} | {snip}"
            if re.search(r"sample", snip, re.I):
                note = "⚠️ 可能是示意值，非即時 | " + note
            return check_range(float(m.group(1)), 0, 100), note
        last_err = "找不到數值 (可能是動態載入或改版)"
    raise FetchError(last_err)


def src_aaii():
    r = get("https://www.aaii.com/sentimentsurvey")
    mb, sb = search([r"Bullish[^0-9%]{0,120}?(\d{1,2}(?:\.\d+)?)\s*%"], r.text)
    mr, sr = search([r"Bearish[^0-9%]{0,120}?(\d{1,2}(?:\.\d+)?)\s*%"], r.text)
    if not (mb and mr):
        raise FetchError("找不到多空比例 (可能被擋或改版)")
    bull, bear = float(mb.group(1)), float(mr.group(1))
    check_range(bull, 0, 100)
    check_range(bear, 0, 100)
    return round(bull - bear, 1), f"多 {bull}% / 空 {bear}% (數值=多空差) | {sb}"


def src_naaim():
    page = "https://naaim.org/programs/naaim-exposure-index/"
    r = get(page)
    link = re.search(r'href="([^"]+?\.xlsx[^"]*)"', r.text, re.I)
    if link:
        try:
            v, note = naaim_from_xlsx(requests.compat.urljoin(page, link.group(1)))
            return check_range(v, -200, 200), note
        except FetchError:
            pass
    m, snip = search([r"NAAIM Number[^0-9\-]{0,300}?(-?\d{1,3}\.\d{1,2})"], r.text)
    if not m:
        raise FetchError("找不到數值 (Excel 與網頁表格都失敗)")
    return check_range(float(m.group(1)), -200, 200), "來自網頁表格 | " + snip


def naaim_from_xlsx(url):
    from openpyxl import load_workbook
    r = get(url, {"Accept": "*/*"})
    try:
        ws = load_workbook(io.BytesIO(r.content), read_only=True, data_only=True).active
        rows = list(ws.iter_rows(values_only=True))
    except Exception:
        raise FetchError("Excel 無法解析")
    for i, row in enumerate(rows[:10]):
        cells = [str(c).strip().lower() if c is not None else "" for c in row]
        if any("naaim number" in c for c in cells):
            col = next(j for j, c in enumerate(cells) if "naaim number" in c)
            dcol = next((j for j, c in enumerate(cells) if "date" in c), None)
            best = None
            for data in rows[i + 1:]:
                if col >= len(data) or not isinstance(data[col], (int, float)):
                    continue
                d = data[dcol] if dcol is not None and dcol < len(data) else None
                if best is None:
                    best = (d, data[col])
                elif isinstance(d, datetime) and (not isinstance(best[0], datetime) or d > best[0]):
                    best = (d, data[col])
            if best:
                d = best[0].strftime("%Y-%m-%d") if isinstance(best[0], datetime) else "未知"
                return round(float(best[1]), 2), f"來自 Excel，資料日 {d}"
    raise FetchError("Excel 內找不到 NAAIM Number 欄位")


def src_crypto():
    r = get("https://api.alternative.me/fng/?limit=1", {"Accept": "application/json"})
    try:
        d = r.json()["data"][0]
        v = float(d["value"])
    except Exception:
        raise FetchError("格式不符 (可能改版)")
    return check_range(v, 0, 100), d.get("value_classification", "")


SOURCES = [
    ("CNN", src_cnn),
    ("thetrading.tools (CNN備援)", src_thetrading),
    ("NAAIM", src_naaim),
    ("AAII", src_aaii),
    ("CFGI Stocks", src_cfgi),
    ("SentiSense", src_sentisense),
    ("AltIndex", src_altindex),
    ("Crypto F&G", src_crypto),
]


# ---------- 主程式 ----------

def main():
    now = datetime.now(ZoneInfo("Asia/Taipei")).strftime("%Y-%m-%d %H:%M")
    rows = []
    for name, fn in SOURCES:
        t0 = time.time()
        try:
            value, note = fn()
            rows.append([now, name, "OK", value, note, round(time.time() - t0, 1), ""])
        except FetchError as e:
            rows.append([now, name, "FAIL", "", "", round(time.time() - t0, 1), e.reason])
        except Exception as e:
            rows.append([now, name, "FAIL", "", "", round(time.time() - t0, 1),
                         f"程式錯誤 ({type(e).__name__})"])

    for row in rows:
        print(" | ".join(str(c) for c in row))

    try:
        import gspread
        gc = gspread.service_account_from_dict(json.loads(os.environ["GOOGLE_SERVICE_ACCOUNT"]))
        ws = gc.open_by_key(SHEET_ID).worksheet(TAB)
        if not ws.row_values(1):
            ws.append_row(HEADER, value_input_option="RAW")
        ws.append_rows(rows, value_input_option="RAW")
        print("已寫入 Sheet")
    except Exception as e:
        print(f"寫入 Sheet 失敗: {type(e).__name__}: {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()
